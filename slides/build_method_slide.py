# /// script
# dependencies = ["python-pptx"]
# ///
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("method-summary.pptx")

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x65, 0x73)
ACCENT = RGBColor(0x0B, 0x5C, 0x8A)
CARD_LINE = RGBColor(0xD0, 0xD7, 0xDE)
WARN = RGBColor(0xB4, 0x53, 0x09)
WARN_FILL = RGBColor(0xFF, 0xF7, 0xED)
FONT = "Calibri"

TITLE = "We will compare three MediaPipe pose models on the Arduino UNO Q, first on video, then on a live webcam"

MARGIN = Inches(0.5)
GAP = Inches(0.25)
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
COL_W = (SLIDE_W - 2 * MARGIN - 2 * GAP) / 3

CARDS = [
    ("WHAT IS COMPARED", [
        ("**Pose Landmarker Lite, Full and Heavy.** Same 33 landmarks, different size and accuracy.", "body"),
        ("**Left out on purpose:**", "body"),
        ("The older `mp.solutions.pose`: unsupported since 2023 and has no build for the Board's Python 3.13.", "bullet"),
        ("Arduino's PoseNet brick: not MediaPipe, 17 keypoints, restricted to the VENTUNO Q.", "bullet"),
    ]),
    ("TWO PHASES", [
        ("**1. Video phase.** Every variant runs on the Board over the same five stock Clips: front view, side view, fast movement, far away or partly hidden, empty room. Each pairing runs 3 times.", "body"),
        ("**2. Webcam phase.** Only variants that pass move on to the Logitech Brio 105 at 1280x720: three 60 second sessions each.", "body"),
        ("If no variant passes, that is the finding, and the webcam demo runs Lite only, labelled as-is.", "muted"),
    ]),
    ("WHAT WE MEASURE", [
        ("**Speed:** FPS, average and slowest-5% frame time, CPU, RAM, temperature.", "body"),
        ("**Quality:** PCK@0.2, detection rate, jitter.", "body"),
        ("The first 30 frames of every Run are dropped as warm-up.", "muted"),
    ]),
    ("HOW ACCURACY IS SCORED", [
        ("**Accuracy is measured against Heavy running on the laptop, not against hand-labelled ground truth.**", "body"),
        ("Those Reference landmarks are a stand-in. PCK@0.2 shows how closely a variant agrees with Heavy, so Heavy on the Board scores close to 100% by construction.", "body"),
    ]),
    ("PASS CRITERIA FOR THE WEBCAM", [
        ("On the Board, across all Clips:", "body"),
        ("Average at least **10 FPS**", "bullet"),
        ("Slowest 5% of frames each **150 ms or less**", "bullet"),
        ("Person found in at least **90%** of frames in Clips that contain one", "bullet"),
        ("**No crashes** or out-of-memory Runs", "bullet"),
        ("**We chose these ourselves**, not from a standard. They can change if you name other targets.", "body"),
    ]),
    ("LAPTOP BASELINE", [
        ("All three variants also run on the laptop over the same Clips and are shown next to the Board's results.", "body"),
    ]),
]
WARN_CARD = 3

FEEDBACK = ("**Feedback wanted before we run it.** Are these Pass criteria the right bar? "
            "Is Heavy an acceptable reference? Is anything missing from the comparison?")


def add_runs(paragraph, text, size, color):
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**"):
            run.text, run.font.bold = part[2:-2], True
        elif part.startswith("`"):
            run.text = part[1:-1]
            run.font.name = "Consolas"
        else:
            run.text = part
        run.font.size = Pt(size)
        run.font.color.rgb = color
        if run.font.name is None:
            run.font.name = FONT


def add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    return frame


def add_card(slide, left, top, width, height, label, paragraphs, warn):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WARN_FILL if warn else RGBColor(0xFF, 0xFF, 0xFF)
    shape.line.color.rgb = WARN if warn else CARD_LINE
    shape.line.width = Pt(1.5 if warn else 1)
    shape.shadow.inherit = False

    pad = Inches(0.2)
    frame = add_textbox(slide, left + pad, top + Inches(0.16), width - 2 * pad, height - Inches(0.3))
    frame.vertical_anchor = MSO_ANCHOR.TOP
    first = frame.paragraphs[0]
    add_runs(first, f"**{label}**", 12, WARN if warn else ACCENT)
    first.space_after = Pt(6)
    for text, style in paragraphs:
        p = frame.add_paragraph()
        p.space_after = Pt(5)
        if style == "bullet":
            text = "•	" + text
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.22)))
            pPr.set("indent", str(-Inches(0.22)))
        add_runs(p, text, 13 if style != "muted" else 12, MUTED if style == "muted" else INK)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title = add_textbox(slide, MARGIN, Inches(0.35), SLIDE_W - 2 * MARGIN, Inches(0.95))
    add_runs(title.paragraphs[0], f"**{TITLE}**", 28, INK)

    card_h = Inches(2.68)
    top0 = Inches(1.35)
    for i, (label, paras) in enumerate(CARDS):
        row, col = divmod(i, 3)
        add_card(slide,
                 MARGIN + col * (COL_W + GAP),
                 top0 + row * (card_h + GAP),
                 COL_W, card_h, label, paras, warn=(i == WARN_CARD))

    feedback = add_textbox(slide, MARGIN, Inches(7.08), SLIDE_W - 2 * MARGIN, Inches(0.3))
    add_runs(feedback.paragraphs[0], FEEDBACK, 14, ACCENT)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "Early method check, sent before any full measurements run. "
        "Design reasoning is in the repo: ADR 0002 (why only the three Pose Landmarker variants), "
        "ADR 0003 (accuracy against Heavy as Reference landmarks, not ground truth), "
        "ADR 0004 (self-chosen Pass criteria). "
        "The natural next step for true accuracy is a labelled dataset such as Penn Action."
    )
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
