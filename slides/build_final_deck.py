"""Build the final mentor deck (issue #11) from the Run files.

Every number on a slide is computed here through bench.report and bench.analysis, never typed by hand.
Run from the repo root:

    $env:PYTHONPATH="src"; uv run --with python-pptx --with matplotlib python slides\\build_final_deck.py

The output embeds the Demo videos from report/demo/ (face-blurred webcam footage), so it is gitignored.
"""
import argparse
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from bench.analysis import RunSummary, analyze, analyze_webcam
from bench.models import VARIANTS
from bench.quality import clip_stem
from bench.report import (
    BOARD,
    LAPTOP,
    RULES,
    by_variant,
    campaign_runs,
    pooled_detection_rate,
    pooled_fps,
    pooled_live,
    pooled_p95_ms,
    pooled_pck,
    provenance,
    spread,
    verdicts,
    webcam_sessions,
)

ROOT = Path(__file__).resolve().parent.parent
VIDEO_CAMPAIGN = "runs/board/video-20261031"
WEBCAM_CAMPAIGN = "runs/board/webcam-20261031T030406Z"
EXTRA_WEBCAM_CAMPAIGN = "runs/board/webcam-20261031T025448Z"
LAPTOP_RUNS = "runs/dev"
DEMO_RUN = "runs/board/video-20261031/video__lite__front__r1__board.jsonl"
DEMO_DIR = ROOT / "report" / "demo"
CHART_DIR = ROOT / "report" / "deck"
OUT = Path(__file__).with_name("final-deck.pptx")

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x65, 0x73)
FAINT = RGBColor(0x9A, 0xA5, 0xB1)
ACCENT = RGBColor(0x0B, 0x5C, 0x8A)
ACCENT_FILL = RGBColor(0xE8, 0xF1, 0xF7)
LINE = RGBColor(0xD0, 0xD7, 0xDE)
WARN = RGBColor(0xB4, 0x53, 0x09)
WARN_FILL = RGBColor(0xFF, 0xF7, 0xED)
FAIL = RGBColor(0xB4, 0x23, 0x18)
PASS = RGBColor(0x1A, 0x7F, 0x37)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
BODY_TOP = Inches(1.75)
CONTENT_W = SLIDE_W - 2 * MARGIN

FPS_BAR, P95_BAR, DETECT_BAR = (rule.limit for rule in RULES[:3])


# ---------------------------------------------------------------- numbers


@dataclass(frozen=True)
class VariantNumbers:
    variant: str
    fps: float
    p95_ms: float
    median_ms: float
    detection: float
    pck: float
    laptop_fps: float
    laptop_pck: float
    passed: bool
    fps_without_empty: float
    fps_empty: float
    max_temp_c: float
    peak_rss_mb: int
    false_detection: float
    jitter: float

    @property
    def name(self) -> str:
        return self.variant.capitalize()


@dataclass(frozen=True)
class Numbers:
    variants: dict[str, VariantNumbers]
    board_runs: dict[str, list[RunSummary]]
    summaries: list[RunSummary]
    webcam: object
    webcam_sessions: list
    webcam_extra: list
    demo_run: RunSummary
    made_by: str
    made_by_extra: str


def compute() -> Numbers:
    paths = [VIDEO_CAMPAIGN, WEBCAM_CAMPAIGN, LAPTOP_RUNS]
    summaries = analyze([ROOT / p for p in paths])
    sessions = analyze_webcam([ROOT / WEBCAM_CAMPAIGN])
    extra = analyze_webcam([ROOT / EXTRA_WEBCAM_CAMPAIGN])
    board, laptop = (by_variant(campaign_runs(summaries, m)) for m in (BOARD, LAPTOP))
    passed = {v.variant: v.passed for v in verdicts(summaries)}
    numbers = {}
    for v in VARIANTS:
        runs = board[v]
        with_person = [r for r in runs if clip_stem(r.clip) != "empty"]
        empty = [r for r in runs if clip_stem(r.clip) == "empty"]
        numbers[v] = VariantNumbers(
            variant=v,
            fps=pooled_fps(runs),
            p95_ms=pooled_p95_ms(runs),
            median_ms=float(np.median([ms for r in runs for ms in r.frame_ms])),
            detection=pooled_detection_rate(runs),
            pck=pooled_pck(runs),
            laptop_fps=pooled_fps(laptop[v]),
            laptop_pck=pooled_pck(laptop[v]),
            passed=passed[v],
            fps_without_empty=pooled_fps(with_person),
            fps_empty=pooled_fps(empty),
            max_temp_c=max(r.max_temp_c for r in runs if r.max_temp_c is not None),
            peak_rss_mb=max(r.peak_rss_mb for r in runs if r.peak_rss_mb is not None),
            false_detection=float(np.mean([r.false_detection_rate for r in empty])),
            jitter=float(np.mean([r.jitter for r in runs if clip_stem(r.clip) == "front"])),
        )
    live = webcam_sessions(sessions)
    demo = next(s for s in summaries if (ROOT / DEMO_RUN).name.startswith(
        f"video__{s.variant}__{clip_stem(s.clip)}__r{s.repeat}__{s.machine}"))
    return Numbers(
        variants=numbers,
        board_runs=board,
        summaries=summaries,
        webcam=pooled_live(live["lite"]),
        webcam_sessions=sorted(live["lite"], key=lambda s: s.session),
        webcam_extra=sorted(webcam_sessions(extra)["lite"], key=lambda s: s.session),
        demo_run=demo,
        made_by=provenance(paths),
        made_by_extra=provenance([EXTRA_WEBCAM_CAMPAIGN]),
    )


# ---------------------------------------------------------------- drawing helpers


def add_text(slide, left, top, width, height, paragraphs, anchor=MSO_ANCHOR.TOP):
    """paragraphs: list of (text, size, color, bold) or (text, size, color, bold, align)."""
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    for i, spec in enumerate(paragraphs):
        text, size, color, bold, *rest = spec
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        if rest:
            p.alignment = rest[0]
        if i:
            p.space_before = Pt(size * 0.55)
        run = p.add_run()
        run.text = text
        run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
        run.font.color.rgb = color
    return box


def add_bullets(slide, left, top, width, height, items, size=18, color=INK, gap=0.7):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for i, item in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_before = Pt(0 if i == 0 else size * gap)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.28)))
        pPr.set("indent", str(-Inches(0.28)))
        bullet = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bullet)
        head, _, rest = item.partition("|")
        run = p.add_run()
        run.text = head
        run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bool(rest)
        run.font.color.rgb = color
        if rest:
            run = p.add_run()
            run.text = rest
            run.font.name, run.font.size = FONT, Pt(size)
            run.font.color.rgb = MUTED
    return box


def rect(slide, left, top, width, height, fill=None, line=None, line_w=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    return shape


def hline(slide, x1, x2, y, color, width=1.5, dash=True):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y, x2, y)
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if dash:
        ln = line.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    return line


def vline(slide, x, y1, y2, color, width=1.5):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y1, x, y2)
    line.line.color.rgb = color
    line.line.width = Pt(width)
    ln = line.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    return line


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        self.count = 0

    def slide(self, headline, sub=None, notes="", numbered=True, section=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.count += 1
        if section:
            add_text(slide, MARGIN, Inches(0.38), CONTENT_W, Inches(0.3), [(section.upper(), 11, ACCENT, True)])
        add_text(slide, MARGIN, Inches(0.68), CONTENT_W, Inches(0.9), [(headline, 28, INK, True)],
                 anchor=MSO_ANCHOR.TOP)
        if sub:
            add_text(slide, MARGIN, Inches(1.62) if len(headline) > 70 else Inches(1.22), CONTENT_W, Inches(0.4),
                     [(sub, 16, MUTED, False)])
        if numbered:
            add_text(slide, SLIDE_W - MARGIN - Inches(1), SLIDE_H - Inches(0.45), Inches(1), Inches(0.25),
                     [(str(self.count), 11, FAINT, False, PP_ALIGN.RIGHT)])
        slide.notes_slide.notes_text_frame.text = notes
        return slide

    def save(self, path):
        self.prs.save(path)


def style_chart(chart, size=14):
    chart.font.name, chart.font.size = FONT, Pt(size)
    chart.font.color.rgb = INK


def manual_plot_area(chart, x, y, w, h):
    """Pin the inner plot area to fractions of the chart frame, so lines drawn over the chart land on its values."""
    plot_area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    for old in plot_area.findall(qn("c:layout")):
        plot_area.remove(old)
    layout = plot_area.makeelement(qn("c:layout"), {})
    manual = layout.makeelement(qn("c:manualLayout"), {})
    for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"),
                     ("x", x), ("y", y), ("w", w), ("h", h)):
        el = manual.makeelement(qn(f"c:{tag}"), {"val": str(val)})
        manual.append(el)
    layout.append(manual)
    plot_area.insert(0, layout)


def hide_gridlines(axis):
    axis.has_major_gridlines = False
    axis.has_minor_gridlines = False


def axis_line(axis, color=LINE):
    axis.format.line.color.rgb = color


# ---------------------------------------------------------------- charts


def fps_bar_chart(slide, n: Numbers, left, top, width, height, machines=(BOARD,), axis_max=None, pass_bar=True):
    """Horizontal bars, one group per variant; a dashed line at the 10 FPS bar is drawn over the plot."""
    data = CategoryChartData()
    order = list(reversed(VARIANTS))
    data.categories = [v.capitalize() for v in order]
    series_names = {BOARD: "Board", LAPTOP: "Laptop"}
    for m in reversed(machines):
        data.add_series(series_names[m],
                        [n.variants[v].fps if m == BOARD else n.variants[v].laptop_fps for v in order])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, left, top, width, height, data)
    chart = frame.chart
    style_chart(chart, 16)
    chart.has_title = False
    shown = [v.fps if m == BOARD else v.laptop_fps for v in n.variants.values() for m in machines]
    top_value = axis_max or max(*shown, FPS_BAR) * 1.25
    value_axis = chart.value_axis
    value_axis.minimum_scale, value_axis.maximum_scale = 0, top_value
    value_axis.visible = False
    hide_gridlines(value_axis)
    cat = chart.category_axis
    cat.tick_labels.font.size = Pt(18)
    cat.tick_labels.font.bold = True
    axis_line(cat)
    cat.format.line.fill.background()
    plot = chart.plots[0]
    plot.gap_width = 45 if len(machines) == 1 else 60
    plot.overlap = 0
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format, labels.number_format_is_linked = '0.0" FPS"', False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(16)
    labels.font.bold = True
    for series in plot.series:
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = ACCENT if series.name == "Board" else FAINT
    chart.has_legend = len(machines) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(14)
    px, py, pw, ph = 0.12, 0.1 if len(machines) > 1 else 0.04, 0.84, 0.86 if len(machines) > 1 else 0.92
    manual_plot_area(chart, px, py, pw, ph)
    if not pass_bar:
        return
    x = left + int(width * (px + pw * FPS_BAR / top_value))
    y1, y2 = top + int(height * py) - Inches(0.05), top + int(height * (py + ph))
    vline(slide, x, y1, y2, FAIL)
    add_text(slide, x - Inches(1.1), y2 + Inches(0.06), Inches(2.2), Inches(0.3),
             [(f"Pass bar: {FPS_BAR:g} FPS", 13, FAIL, True, PP_ALIGN.CENTER)])


def session_chart(slide, n: Numbers, left, top, width, height):
    """Processed FPS per webcam Session, with the Video phase speed of the same model as a dashed line."""
    data = CategoryChartData()
    data.categories = [f"Session {x.session}" for x in n.webcam_sessions]
    data.add_series("Processed FPS", [x.stats.processed_fps for x in n.webcam_sessions])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, data).chart
    style_chart(chart, 13)
    chart.has_title = False
    chart.has_legend = False
    video_fps = n.variants["lite"].fps
    top_value = max(*(x.stats.processed_fps for x in n.webcam_sessions), video_fps) * 1.15
    va = chart.value_axis
    va.minimum_scale, va.maximum_scale = 0, top_value
    va.visible = False
    hide_gridlines(va)
    ca = chart.category_axis
    ca.tick_labels.font.size = Pt(13)
    ca.tick_labels.font.color.rgb = MUTED
    axis_line(ca)
    plot = chart.plots[0]
    plot.gap_width = 80
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format, labels.number_format_is_linked = '0.0" FPS"', False
    labels.position = XL_LABEL_POSITION.INSIDE_END
    labels.font.size, labels.font.bold = Pt(14), True
    labels.font.color.rgb = WHITE
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = ACCENT
    px, py, pw, ph = 0.02, 0.04, 0.78, 0.84
    manual_plot_area(chart, px, py, pw, ph)
    y = top + int(height * (py + ph * (1 - video_fps / top_value)))
    x_end = left + int(width * (px + pw))
    hline(slide, left + int(width * px), x_end + Inches(0.1), y, MUTED, 1.25)
    add_text(slide, x_end + Inches(0.18), y - Inches(0.25), Inches(1.5), Inches(0.5),
             [("Video phase", 12, MUTED, True), (f"{video_fps:.1f} FPS", 12, MUTED, True)])


def frame_time_png(n: Numbers, path: Path) -> Path:
    plt.rcParams.update({"font.family": "Calibri", "font.size": 15, "axes.spines.top": False,
                         "axes.spines.right": False, "axes.edgecolor": "#D0D7DE", "xtick.color": "#5B6573",
                         "ytick.color": "#5B6573"})
    fig, ax = plt.subplots(figsize=(7.4, 4.6))
    data = [[ms for r in n.board_runs[v] for ms in r.frame_ms] for v in VARIANTS]
    ax.boxplot(data, whis=(5, 95), showfliers=False, patch_artist=True, widths=0.45,
               boxprops={"facecolor": "#E8F1F7", "edgecolor": "#0B5C8A"},
               whiskerprops={"color": "#0B5C8A"}, capprops={"color": "#0B5C8A"},
               medianprops={"color": "#0B5C8A", "linewidth": 2.5})
    ax.set_xticks(range(1, len(VARIANTS) + 1), [v.capitalize() for v in VARIANTS], fontsize=17,
                  fontweight="bold", color="#1F2937")
    ax.axhline(P95_BAR, color="#B42318", linestyle="--", linewidth=1.6)
    ax.text(0.55, P95_BAR - 45, f"Pass bar: p95 ≤ {P95_BAR:g} ms", color="#B42318", ha="left",
            fontsize=13, fontweight="bold")
    ax.set_ylabel("time per frame (ms)", color="#5B6573")
    ax.set_ylim(0, None)
    ax.yaxis.grid(True, color="#EEF1F4")
    ax.set_axisbelow(True)
    fig.tight_layout()
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=220)
    plt.close(fig)
    return path


def scatter_chart(slide, n: Numbers, left, top, width, height):
    data = XyChartData()
    for v in VARIANTS:
        s = data.add_series(v.capitalize())
        s.add_data_point(n.variants[v].fps, n.variants[v].pck)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, left, top, width, height, data)
    chart = frame.chart
    style_chart(chart, 14)
    chart.has_title = False
    chart.has_legend = False
    va, ca = chart.value_axis, chart.category_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 80, 100, 5
    ca.minimum_scale, ca.maximum_scale, ca.major_unit = 0, 6, 1
    for axis, title in ((va, "agreement with Heavy (PCK@0.2, %)"), (ca, "speed on the Board (FPS)")):
        axis.has_title = True
        axis.axis_title.text_frame.text = title
        run = axis.axis_title.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.bold, run.font.name = Pt(14), False, FONT
        run.font.color.rgb = MUTED
        axis.tick_labels.font.size = Pt(13)
        axis.tick_labels.font.color.rgb = MUTED
        axis_line(axis)
        axis.major_gridlines.format.line.color.rgb = RGBColor(0xEE, 0xF1, 0xF4)
    ca.has_major_gridlines = False
    for series, v in zip(chart.plots[0].series, VARIANTS):
        series.marker.style = XL_MARKER_STYLE.CIRCLE
        series.marker.size = 16
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = ACCENT if v == "lite" else FAINT
        series.marker.format.line.fill.background()
        series.format.line.fill.background()
        point = series.points[0]
        point.data_label.has_text_frame = True
        point.data_label.text_frame.text = v.capitalize()
        point.data_label.position = XL_LABEL_POSITION.RIGHT
        r = point.data_label.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.name = Pt(16), True, FONT
        r.font.color.rgb = ACCENT if v == "lite" else INK


def temperature_chart(slide, n: Numbers, left, top, width, height):
    data = XyChartData()
    colors = {"lite": ACCENT, "full": MUTED, "heavy": WARN}
    for v in VARIANTS:
        run = max(n.board_runs[v], key=lambda r: max((t for _, t in r.temps), default=0))
        s = data.add_series(f"{v.capitalize()} ({clip_stem(run.clip)} Clip)")
        t0 = run.temps[0][0]
        for t, c in run.temps:
            s.add_data_point((t - t0) / 1000 / 60, c)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, left, top, width, height, data)
    chart = frame.chart
    style_chart(chart, 12)
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(12)
    va, ca = chart.value_axis, chart.category_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 30, 60, 10
    ca.minimum_scale, ca.major_unit = 0, 2
    for axis, title in ((va, "CPU °C"), (ca, "minutes into the hottest Run")):
        axis.has_title = True
        axis.axis_title.text_frame.text = title
        r = axis.axis_title.text_frame.paragraphs[0].runs[0]
        r.font.size, r.font.bold, r.font.name = Pt(12), False, FONT
        r.font.color.rgb = MUTED
        axis.tick_labels.font.size = Pt(11)
        axis.tick_labels.font.color.rgb = MUTED
        axis_line(axis)
    va.major_gridlines.format.line.color.rgb = RGBColor(0xEE, 0xF1, 0xF4)
    ca.has_major_gridlines = False
    for series, v in zip(chart.plots[0].series, VARIANTS):
        series.smooth = False
        series.format.line.color.rgb = colors[v]
        series.format.line.width = Pt(1.5)


def table(slide, left, top, width, rows, col_widths, row_h=Inches(0.55), size=16, header_size=13,
          highlight_row=None, cell_colors=None, left_cols=1):
    """rows[0] is the header. cell_colors maps (row, col) to a text colour."""
    shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, row_h * len(rows))
    tbl = shape.table
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tbl_pr.set(attr, "0")
    style = tbl_pr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid; borders set below
    for i, w in enumerate(col_widths):
        tbl.columns[i].width = Emu(int(width * w / sum(col_widths)))
    for r, row in enumerate(rows):
        tbl.rows[r].height = row_h
        for c, text in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT_FILL if r == highlight_row else WHITE
            frame = cell.text_frame
            frame.word_wrap = True
            lines = text.split("\n")
            for i, line in enumerate(lines):
                p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                p.alignment = PP_ALIGN.LEFT if c < left_cols else PP_ALIGN.RIGHT
                run = p.add_run()
                run.text = line
                run.font.name = FONT
                if r == 0:
                    run.font.size = Pt(header_size if i == 0 else header_size - 2)
                    run.font.bold = i == 0
                    run.font.color.rgb = MUTED
                else:
                    run.font.size = Pt(size)
                    run.font.bold = c == 0 or r == highlight_row
                    run.font.color.rgb = (cell_colors or {}).get((r, c), INK)
            borders(cell, bottom=INK if r == 0 else LINE)
    return shape


def borders(cell, bottom):
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tc_pr.findall(qn(tag)):
            tc_pr.remove(old)
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tc_pr.makeelement(qn(tag), {"w": str(Pt(1 if tag == "a:lnB" else 0)), "cmpd": "sng"})
        if tag == "a:lnB":
            fill = ln.makeelement(qn("a:solidFill"), {})
            fill.append(fill.makeelement(qn("a:srgbClr"), {"val": str(bottom)}))
        else:
            fill = ln.makeelement(qn("a:noFill"), {})
        ln.append(fill)
        tc_pr.append(ln)


def stat_tile(slide, left, top, width, height, value, label, verdict=None, accent=ACCENT):
    rect(slide, left, top, width, height, fill=WHITE, line=LINE)
    rect(slide, left, top, Inches(0.07), height, fill=accent)
    pad = Inches(0.28)
    paras = [(value, 34, accent, True), (label, 14, MUTED, False)]
    if verdict:
        paras.append(verdict)
    add_text(slide, left + pad, top + Inches(0.2), width - pad - Inches(0.15), height - Inches(0.3), paras)


def video(slide, path: Path, poster: Path, left, top, width):
    height = int(width * 9 / 16)
    rect(slide, left - Inches(0.02), top - Inches(0.02), width + Inches(0.04), height + Inches(0.04), line=LINE)
    slide.shapes.add_movie(str(path), left, top, width, height, poster_frame_image=str(poster), mime_type="video/mp4")
    return height


# ---------------------------------------------------------------- slides


def pct(x: float) -> str:
    return f"{x:.0f}%"


def build(n: Numbers, out: Path) -> Path:
    d = Deck()
    lite, full, heavy = (n.variants[v] for v in VARIANTS)
    w = n.webcam
    best_laptop_ratio = min(v.laptop_fps / v.fps for v in n.variants.values())
    worst_laptop_ratio = max(v.laptop_fps / v.fps for v in n.variants.values())
    max_temp = max(v.max_temp_c for v in n.variants.values())
    provenance_note = f"Numbers: {n.made_by}. Built by slides/build_final_deck.py."

    # Title --------------------------------------------------------------------------------------------
    s = d.slide("", numbered=False, notes=(
        "Final results of the MediaPipe pose benchmark on the Arduino UNO Q. "
        "Answer first, then the evidence, then the limits and next steps. About 12 minutes."))
    rect(s, 0, 0, Inches(0.18), SLIDE_H, fill=ACCENT)
    add_text(s, Inches(1.1), Inches(2.3), Inches(11.6), Inches(3.2), [
        ("BENCHMARK RESULTS", 14, ACCENT, True),
        ("Testing MediaPipe pose models on the Arduino UNO Q", 36, INK, True),
        ("Which model runs best on the Board, and how fast: results of the Video phase and Webcam phase", 22, MUTED, False),
    ])
    add_text(s, Inches(1.1), Inches(6.3), Inches(11), Inches(0.4), [("November 2026", 14, FAINT, False)])

    # 1 Recommendation ---------------------------------------------------------------------------------
    s = d.slide(
        "Use Lite. No model reached real-time on the Board, and Lite came closest",
        sub=f"It reaches {pct(100 * lite.fps / FPS_BAR)} of the {FPS_BAR:g} FPS bar, so it fits uses where about "
            f"{lite.fps:.0f} pose updates a second is enough.",
        section="Recommendation",
        notes=(
            f"All three Model variants failed our Pass criteria, and all on speed only. Detection passed everywhere. "
            f"Lite is the fastest ({lite.fps:.1f} FPS on the Board) and the only one already shown live on the Board. "
            f"Full costs {pct(100 * (1 - full.fps / lite.fps))} of Lite's speed for {full.pck - lite.pck:.1f} more "
            f"points of agreement with Heavy. Heavy runs at {heavy.fps:.1f} FPS. "
            f"So: Lite, used where about {lite.fps:.0f} pose updates a second is enough, or with acceleration (Next steps). "
            + provenance_note))
    tile_w, tile_h, gap = Inches(3.86), Inches(1.85), Inches(0.25)
    top = Inches(2.35)
    tiles = [
        (f"{lite.fps:.1f} FPS", "average speed on the Board", (f"Pass bar {FPS_BAR:g} FPS: not met", 13, FAIL, True)),
        (f"{lite.p95_ms:.0f} ms", "p95 frame time (95% of frames are faster)",
         (f"Pass bar {P95_BAR:g} ms: not met", 13, FAIL, True)),
        (f"{lite.detection:.0f}%", "of frames with a person found",
         (f"Pass bar {DETECT_BAR:g}%: met", 13, PASS, True)),
    ]
    for i, (value, label, verdict) in enumerate(tiles):
        stat_tile(s, MARGIN + i * (tile_w + gap), top, tile_w, tile_h, value, label, verdict)
    add_bullets(s, MARGIN, Inches(4.75), CONTENT_W, Inches(1.6), [
        f"Full|: {pct(100 * (1 - full.fps / lite.fps))} slower than Lite, for {full.pck - lite.pck:.1f} points more "
        "agreement with Heavy",
        f"Heavy|: {heavy.fps:.1f} FPS, {lite.fps / heavy.fps:.1f}× slower than Lite",
    ], size=18, gap=0.8)

    # 2 Method -----------------------------------------------------------------------------------------
    s = d.slide(
        "All three models ran on video on the Board; Lite alone went on to a live webcam",
        section="Method",
        notes=(
            "Same method as the slide sent before the measurements, now carried out. "
            "Runs were shuffled with a recorded seed and the Board cooled between Runs. "
            "The first 30 frames of each video Run and the first 5 s of each webcam Session are warm-up and not counted. "
            "Software: mediapipe 0.10.18 on Python 3.12 on both machines (ADR 0005: the 1.x wheels crash on the "
            "Board's Cortex-A53 cores). Accuracy is measured against Heavy on the laptop (ADR 0003), not ground truth. "
            "Pass criteria: ADR 0004."))
    col_w = (CONTENT_W - 2 * Inches(0.35)) / 3
    columns = [
        ("WHAT RAN", [
            "Pose Landmarker Lite, Full, Heavy|: same 33 landmarks, different size",
            "On the Board's CPU|: Arduino UNO Q 4 GB, Debian",
            "mediapipe 0.10.18|, Python 3.12",
        ]),
        ("TWO PHASES", [
            f"Video phase|: 5 Clips × 3 repeats per model, {len(campaign_runs(n.summaries, BOARD))} Runs",
            f"Webcam phase|: Lite, {len(n.webcam_sessions)} × 60 s live from a Logitech Brio 105",
            "Laptop baseline|: the same Clips on the laptop",
        ]),
        ("WHAT COUNTS", [
            f"Pass bar (ours)|: ≥ {FPS_BAR:g} FPS, p95 frame time ≤ {P95_BAR:g} ms, person found ≥ "
            f"{DETECT_BAR:g}%, no failed Runs",
            "Agreement with Heavy|: PCK@0.2 against Heavy on the laptop, a stand-in for hand-labelled truth",
        ]),
    ]
    for i, (label, items) in enumerate(columns):
        left = MARGIN + i * (col_w + Inches(0.35))
        rect(s, left, Inches(1.85), col_w, Inches(0.05), fill=ACCENT)
        add_text(s, left, Inches(2.05), col_w, Inches(0.4), [(label, 13, ACCENT, True)])
        add_bullets(s, left, Inches(2.6), col_w, Inches(4), items, size=18, gap=0.9)

    # 3 Summary table ----------------------------------------------------------------------------------
    s = d.slide(
        f"All three models failed on speed alone; each found the person in "
        f"{min(v.detection for v in n.variants.values()):.0f}% of frames or more",
        section="Results",
        notes=(
            "Board and laptop numbers pool every frame of the 15 Runs per model (5 Clips × 3 repeats), after warm-up. "
            "Agreement with Heavy is PCK@0.2 against Heavy on the laptop, so Heavy is high partly by construction. "
            "False detections on the empty-room Clip (frames with a person reported): "
            + ", ".join(f"{v.name} {v.false_detection:.1f}%" for v in (lite, full, heavy))
            + ". Jitter on the front Clip, in torso lengths: "
            + ", ".join(f"{v.name} {v.jitter:.4f}" for v in (lite, full, heavy)) + ". "
            f"Full is {pct(100 * (1 - full.fps / lite.fps))} slower than Lite; Heavy is {lite.fps / heavy.fps:.1f}× "
            f"slower. On the empty-room Clip all three run at about {lite.fps_empty:.0f} FPS "
            f"(Lite {lite.fps_empty:.1f}, Full {full.fps_empty:.1f}, Heavy {heavy.fps_empty:.1f}): with no person "
            "found, only the person detector runs, not the pose model. Without that Clip: "
            f"Lite {lite.fps_without_empty:.1f}, Full {full.fps_without_empty:.1f}, "
            f"Heavy {heavy.fps_without_empty:.1f} FPS. "
            + provenance_note))
    hdr = ["", f"FPS\nbar ≥ {FPS_BAR:g}", f"p95 frame time\nbar ≤ {P95_BAR:g} ms", f"Person found\nbar ≥ {DETECT_BAR:g}%",
           "Agreement\nwith Heavy", "Laptop\nFPS", "Verdict"]
    rows, colors = [hdr], {}
    for r, v in enumerate((lite, full, heavy), start=1):
        rows.append([v.name, f"{v.fps:.1f}", f"{v.p95_ms:.0f} ms", f"{v.detection:.1f}%", f"{v.pck:.1f}%",
                     f"{v.laptop_fps:.1f}", "PASS" if v.passed else "FAIL"])
        colors[(r, 1)] = INK if v.fps >= FPS_BAR else FAIL
        colors[(r, 2)] = INK if v.p95_ms <= P95_BAR else FAIL
        colors[(r, 6)] = PASS if v.passed else FAIL
    table(s, MARGIN, Inches(2.1), CONTENT_W, rows, [1.6, 1.5, 1.8, 1.8, 1.7, 1.4, 1.3], row_h=Inches(0.78),
          size=22, header_size=14, cell_colors=colors)
    add_text(s, MARGIN, Inches(5.5), CONTENT_W, Inches(0.8), [
        ("Red: misses the Pass bar. Speed numbers are the Board's; the laptop column is for comparison only.",
         14, MUTED, False)])

    # 4b frame times -----------------------------------------------------------------------------------
    png = frame_time_png(n, CHART_DIR / "frame-times.png")
    s = d.slide(
        f"The slowness is steady, not occasional: Lite's typical frame takes {lite.median_ms:.0f} ms",
        sub=f"The p95 bar of {P95_BAR:g} ms fails because every frame is slow, not because of a few spikes.",
        section="Video phase",
        notes=(
            "Box: middle half of frames. Whiskers: 5th to 95th percentile. "
            f"Median frame time: Lite {lite.median_ms:.0f} ms, Full {full.median_ms:.0f} ms, Heavy {heavy.median_ms:.0f} ms. "
            f"95th percentile (p95): Lite {lite.p95_ms:.0f}, Full {full.p95_ms:.0f}, Heavy {heavy.p95_ms:.0f} ms. "
            "Heavy's low whisker is the empty-room Clip, where only the person detector runs. " + provenance_note))
    s.shapes.add_picture(str(png), MARGIN + Inches(1.2), Inches(2.0), height=Inches(5.0))

    # 4c accuracy vs speed -----------------------------------------------------------------------------
    s = d.slide(
        f"Full agrees with Heavy {full.pck - lite.pck:.1f} points more, at {pct(100 * (1 - full.fps / lite.fps))} "
        "less speed; Lite is the better trade",
        section="Video phase",
        notes=(
            f"Agreement with Heavy on the laptop (PCK@0.2): Lite {lite.pck:.1f}%, Full {full.pck:.1f}%, "
            f"Heavy {heavy.pck:.1f}%. Heavy's score is high partly by construction: it is compared with itself, "
            "run on a different machine. Board and laptop scores are close: "
            f"Lite {lite.pck:.1f}% on the Board vs {lite.laptop_pck:.1f}% on the laptop. " + provenance_note))
    scatter_chart(s, n, MARGIN, Inches(1.95), Inches(8.2), Inches(5.0))
    add_text(s, MARGIN + Inches(8.7), Inches(2.4), Inches(3.4), Inches(3), [
        ("Heavy scores high partly because it is the reference it is scored against.", 17, MUTED, False),
        ("Higher means closer to Heavy, not proven correct.", 17, MUTED, False),
    ])

    # 4d Demo video: Clip ------------------------------------------------------------------------------
    demo = n.demo_run
    s = d.slide(
        "On a Clip, Lite tracks the person well, just slowly",
        section="Video phase · Demo video",
        notes=(
            f"Demo video: the Board's own landmarks from one Video phase Run (Lite, front Clip, repeat {demo.repeat}), "
            f"drawn on the same frames. The Board ran this Run at {demo.fps:.1f} FPS; the video plays at the Clip's "
            "30 fps, so it looks smoother than it was live. Not a measurement. Run file: " + DEMO_RUN))
    vw = Inches(8.9)
    video(s, DEMO_DIR / "video-phase.mp4", DEMO_DIR / "video-phase.png", MARGIN, Inches(1.6), vw)
    add_bullets(s, MARGIN + vw + Inches(0.4), Inches(1.7), SLIDE_W - 2 * MARGIN - vw - Inches(0.4), Inches(4), [
        "Lite on the Board|, front Clip",
        f"Ran at {demo.fps:.1f} FPS|; played back at 30 fps",
        "Skeleton = the Board's own output",
    ], size=16, gap=1.0)

    # 5a Webcam ----------------------------------------------------------------------------------------
    s = d.slide(
        f"Live, Lite updates the pose about {w.processed_fps:.0f} times a second, "
        f"about {w.delay_avg_ms / 1000:.1f} s behind",
        sub="One frame goes to MediaPipe at a time; camera frames that arrive meanwhile are dropped (ADR 0006).",
        section="Webcam phase",
        notes=(
            f"Lite, labelled 'as-is' because it failed the Pass criteria (ADR 0004). {len(n.webcam_sessions)} Sessions "
            f"of 60 s, first 5 s warm-up. Per Session: "
            + ", ".join(f"{x.stats.processed_fps:.1f} FPS / {x.stats.delay_avg_ms:.0f} ms" for x in n.webcam_sessions)
            + f". Delay runs from capture to result; p95 {w.delay_p95_ms:.0f} ms. About 48 ms of it is USB transfer and "
            "MJPG decode, inferred from per-frame timings, not profiled. MediaPipe 0.10.18 queues live frames, so the "
            "Session sends one frame only when the last result is back (ADR 0006); dropped = camera frames never sent. "
            f"An extra Campaign ({EXTRA_WEBCAM_CAMPAIGN}) with some still Sessions gave the same speed: "
            + ", ".join(f"{x.stats.processed_fps:.1f}" for x in n.webcam_extra)
            + f" FPS. {n.made_by_extra}. " + provenance_note))
    tiles = [
        (f"{w.processed_fps:.1f} FPS", f"processed, of {w.camera_fps:.0f} from the camera"),
        (f"{w.dropped_pct:.0f}%", "of camera frames dropped"),
        (f"{w.delay_avg_ms:.0f} ms", f"capture to result (p95 {w.delay_p95_ms:.0f} ms)"),
        (f"{w.detection_rate:.0f}%", "of processed frames with a person found"),
    ]
    tile_w = (CONTENT_W - 3 * Inches(0.25)) / 4
    for i, (value, label) in enumerate(tiles):
        stat_tile(s, MARGIN + i * (tile_w + Inches(0.25)), Inches(2.05), tile_w, Inches(1.75), value, label)
    session_chart(s, n, MARGIN, Inches(4.2), Inches(7.4), Inches(2.75))
    add_text(s, MARGIN + Inches(7.9), Inches(4.6), CONTENT_W - Inches(7.9), Inches(1.8), [
        ("Each Session landed close to the Video phase speed.", 16, INK, True),
        (f"{len(n.webcam_sessions)} Sessions × 60 s, Brio 105 at 1280×720, 30 fps.", 14, MUTED, False)])

    # 5b Demo video: webcam ----------------------------------------------------------------------------
    s = d.slide(
        "Live, the skeleton follows the person a few times a second",
        section="Webcam phase · Demo video",
        notes=(
            "Demo video from a Demo recording on the Board: every camera frame, with the latest skeleton that had "
            "arrived when the frame was read, as a live viewer would see it. It is not measured: the preview adds "
            "work on the Board. The face is blurred after pose estimation, so the blur does not change the landmarks. "
            "The numbers are on the previous slide."))
    video(s, DEMO_DIR / "webcam.mp4", DEMO_DIR / "webcam.png", MARGIN, Inches(1.6), vw)
    add_bullets(s, MARGIN + vw + Inches(0.4), Inches(1.7), SLIDE_W - 2 * MARGIN - vw - Inches(0.4), Inches(4), [
        "Lite on the Board|, live",
        "Skeleton shown as it arrived|, so it trails movement",
        "Face blurred|; not measured",
    ], size=16, gap=1.0)

    # 6 Laptop baseline --------------------------------------------------------------------------------
    s = d.slide(
        f"The same models run {best_laptop_ratio:.1f}–{worst_laptop_ratio:.1f}× faster on the laptop",
        sub="Lite and Full agree with Heavy about as well on the Board as on the laptop: slower, not worse.",
        section="Laptop baseline",
        notes=(
            "Laptop baseline: the same Clips, the same mediapipe 0.10.18 and Python 3.12, on the Windows laptop's CPU. "
            + " ".join(f"{v.name}: Board {v.fps:.1f} vs laptop {v.laptop_fps:.1f} FPS "
                       f"({v.laptop_fps / v.fps:.1f}×), agreement {v.pck:.1f}% vs {v.laptop_pck:.1f}%."
                       for v in (lite, full, heavy))
            + " Laptop repeats varied more than the Board's (appendix A), since the laptop ran other software. "
            + provenance_note))
    fps_bar_chart(s, n, MARGIN, Inches(2.0), Inches(9.5), Inches(4.9), machines=(BOARD, LAPTOP), pass_bar=False)

    # 7 Left out ---------------------------------------------------------------------------------------
    s = d.slide(
        "Two things were left out on purpose, to keep the comparison fair",
        section="Scope",
        notes=(
            "ADR 0001: plain Python venv over SSH, not App Lab. App Lab runs code in a container next to Arduino's own "
            "OpenCV, which may clash with the one mediapipe installs, adds overhead to the timings, and defaults the "
            "camera to 640×480 at 10 fps. "
            "ADR 0002: only the three Pose Landmarker variants. The legacy mp.solutions.pose is unsupported since 2023 "
            "and missing from mediapipe since 0.10.30. Arduino's pose brick is PoseNet (17 keypoints), not MediaPipe, "
            "and limited to the VENTUNO Q."))
    col_w = (CONTENT_W - Inches(0.5)) / 2
    for i, (label, title, items) in enumerate([
        ("ADR 0001", "Arduino App Lab", [
            "Container overhead|, which would mix into the timings",
            "A second OpenCV|, which may clash with MediaPipe's",
            "Better suited to a final demo|: see Next steps",
        ]),
        ("ADR 0002", "Other pose models", [
            "Legacy mp.solutions.pose|: unsupported since 2023",
            "Arduino's pose brick|: PoseNet, not MediaPipe; VENTUNO Q only",
            "Either can be added later|, without changing this comparison",
        ]),
    ]):
        left = MARGIN + i * (col_w + Inches(0.5))
        rect(s, left, Inches(1.85), col_w, Inches(0.05), fill=ACCENT)
        add_text(s, left, Inches(2.05), col_w, Inches(0.9), [(label, 13, ACCENT, True), (title, 24, INK, True)])
        add_bullets(s, left, Inches(3.2), col_w, Inches(3.5), items, size=18, gap=0.9)

    # 8 Limits -----------------------------------------------------------------------------------------
    s = d.slide(
        "Four limits to keep in mind when reading these results",
        section="Limits",
        notes=(
            f"Heatsink: the Board ran bare. Hottest reading in any Video phase Run: {max_temp:.1f} °C. We did not "
            "check the QRB2210's throttling point, so whether a heatsink would change the speed is unknown. "
            "Pass criteria: chosen by us for an exploratory study, not from a standard (ADR 0004); they can change. "
            "Reference: Heavy on the laptop, not hand-labelled ground truth (ADR 0003). "
            "Unverified: the split of the webcam delay (about 48 ms of camera transfer and decode) is inferred from "
            "per-frame timings, not profiled; in both webcam Campaigns the second Session was the fastest, which is "
            "not explained."))
    add_bullets(s, MARGIN, Inches(1.9), Inches(6.4), Inches(5), [
        "The accuracy score is agreement with Heavy|, not hand-labelled ground truth",
        f"No heatsink|: the Board peaked at {max_temp:.0f} °C; its effect on speed is untested",
        "The Pass bars are ours|, not from a standard",
        "Some details are unverified|: the webcam delay split, why Session 2 ran fastest",
    ], size=18, gap=1.1)
    temperature_chart(s, n, MARGIN + Inches(6.8), Inches(1.9), Inches(5.3), Inches(4.4))
    add_text(s, MARGIN + Inches(6.8), Inches(6.35), Inches(5.3), Inches(0.4),
             [("CPU temperature during the hottest Video phase Run of each model", 12, MUTED, False)])

    # 9 Next steps -------------------------------------------------------------------------------------
    s = d.slide(
        "Next: try to close the speed gap, then measure true accuracy",
        section="Next steps",
        notes=(
            "Ordered by how directly each answers the open question. Acceleration is the most likely way to narrow "
            f"a {FPS_BAR / lite.fps:.1f}× gap, though untested here; a heatsink only matters if the Board throttles. "
            "Penn Action has hand-labelled joints, which would replace the Heavy stand-in. "
            "App Lab is Arduino's own way to ship an app (ADR 0001)."))
    steps = [
        ("1", "LiteRT or GPU acceleration", f"Could narrow the {FPS_BAR / lite.fps:.1f}× speed gap; untested on the UNO Q"),
        ("2", "Heatsink comparison", "Same Campaign with a heatsink, to see if heat limits speed"),
        ("3", "A labelled dataset such as Penn Action", "Accuracy against real ground truth, not Heavy"),
        ("4", "An App Lab demo", "Lite packaged the way Arduino ships apps"),
    ]
    row_h = Inches(1.15)
    for i, (num, title, why) in enumerate(steps):
        top = Inches(1.95) + i * row_h
        add_text(s, MARGIN, top, Inches(0.7), Inches(0.8), [(num, 30, ACCENT, True)])
        add_text(s, MARGIN + Inches(0.8), top + Inches(0.05), Inches(10.5), Inches(0.9),
                 [(title, 22, INK, True), (why, 16, MUTED, False)])
        if i < len(steps) - 1:
            hline(s, MARGIN + Inches(0.8), SLIDE_W - MARGIN, top + row_h - Inches(0.12), LINE, 0.75, dash=False)

    # Appendix A: spread --------------------------------------------------------------------------------
    rows_spread = spread(n.summaries)
    pcts = {m: [x.fps_spread_pct for x in rows_spread if x.machine == m] for m in (BOARD, LAPTOP)}
    spread_note = (f"FPS spread across repeats (range as % of mean): Board {min(pcts[BOARD]):.1f}–"
                   f"{max(pcts[BOARD]):.1f}%, laptop {min(pcts[LAPTOP]):.1f}–{max(pcts[LAPTOP]):.1f}%.")
    s = d.slide(
        "Appendix A: repeats agreed closely on the Board, less so on the laptop",
        section="Appendix",
        notes=("FPS mean and range across the 3 repeats of each model and Clip, and the range of the p95 frame "
               "time. " + spread_note + " " + provenance_note))
    rows = [["Model", "Clip", "Board FPS\nmean (min–max)", "Board p95 frame time\nms (min–max)",
             "Laptop FPS\nmean (min–max)"]]
    by_key = {(x.variant, x.clip, x.machine): x for x in rows_spread}
    for v in VARIANTS:
        for clip in ("front", "side", "fast", "far", "empty"):
            b, lap = by_key[(v, clip, BOARD)], by_key[(v, clip, LAPTOP)]
            rows.append([v.capitalize() if clip == "front" else "", clip,
                         f"{b.fps_mean:.2f} ({b.fps_min:.2f}–{b.fps_max:.2f})",
                         f"{b.p95_min:.0f}–{b.p95_max:.0f}",
                         f"{lap.fps_mean:.1f} ({lap.fps_min:.1f}–{lap.fps_max:.1f})"])
    table(s, MARGIN, Inches(1.4), CONTENT_W, rows, [1.2, 1.2, 2.6, 2.4, 2.6], row_h=Inches(0.33), size=12,
          header_size=12, left_cols=2)

    # Appendix B: webcam -------------------------------------------------------------------------------
    s = d.slide(
        "Appendix B: the three webcam Sessions",
        section="Appendix",
        notes=("Lite (as-is) on the Board, Campaign " + WEBCAM_CAMPAIGN + ". " + provenance_note))
    rows = [["Session", "Camera FPS", "Processed FPS", "Dropped", "Delay avg", "Delay p95", "Person found",
             "Peak RSS", "Max temp"]]
    for x in n.webcam_sessions:
        st = x.stats
        rows.append([str(x.session), f"{st.camera_fps:.1f}", f"{st.processed_fps:.1f}", f"{st.dropped_pct:.1f}%",
                     f"{st.delay_avg_ms:.0f} ms", f"{st.delay_p95_ms:.0f} ms", f"{st.detection_rate:.1f}%",
                     f"{x.peak_rss_mb} MB", f"{x.max_temp_c:.1f} °C"])
    rows.append(["All", f"{w.camera_fps:.1f}", f"{w.processed_fps:.1f}", f"{w.dropped_pct:.1f}%",
                 f"{w.delay_avg_ms:.0f} ms", f"{w.delay_p95_ms:.0f} ms", f"{w.detection_rate:.1f}%",
                 f"{max(x.peak_rss_mb for x in n.webcam_sessions)} MB",
                 f"{max(x.max_temp_c for x in n.webcam_sessions):.1f} °C"])
    table(s, MARGIN, Inches(1.8), CONTENT_W, rows, [1.2, 1.3, 1.5, 1.2, 1.3, 1.3, 1.5, 1.3, 1.3],
          row_h=Inches(0.6), size=16, header_size=13, highlight_row=len(rows) - 1)

    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(out)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the final mentor deck from the Run files.")
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args()
    print(build(compute(), args.out))


if __name__ == "__main__":
    main()
